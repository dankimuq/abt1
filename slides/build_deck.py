"""Build the editable DSN 2027 ABT research presentation."""
from __future__ import annotations

from pathlib import Path
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("ABT_DSN_2027.pptx")
W, H = 13.333, 7.5
NAVY = "0B1220"
PANEL = "131F31"
PANEL_2 = "18283D"
WHITE = "F4F7FB"
MUTED = "AAB7C8"
CYAN = "4DD6E5"
BLUE = "5D8CFF"
AMBER = "FFC66D"
GREEN = "73D7A5"
RED = "FF7B86"
LINE = "2C3D54"
FONT = "Aptos"


def color(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def add_box(slide, x, y, w, h, fill=PANEL, line=None, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.fill.background() if line is None else None
    if line:
        shape.line.color.rgb = color(line)
        shape.line.width = Pt(1)
    if radius:
        shape.adjustments[0] = 0.12
    return shape


def add_text(
    slide, text, x, y, w, h, size=18, fill=WHITE, bold=False,
    font=FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
    margin=0.04, italic=False,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color(fill)
    return shape


def add_bullets(slide, items, x, y, w, h, size=17, fill=WHITE, gap=8):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"•  {item}"
        p.space_after = Pt(gap)
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color(fill)
    return shape


def add_base(prs, section, title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(NAVY)
    add_text(slide, section.upper(), 0.62, 0.35, 4.0, 0.26, 10, CYAN, True)
    add_text(slide, title, 0.62, 0.72, 12.0, 0.58, 28, WHITE, True)
    if subtitle:
        add_text(slide, subtitle, 0.66, 1.33, 12.0, 0.40, 13, MUTED)
    add_box(slide, 0.65, 7.08, 12.02, 0.012, LINE, radius=False)
    add_text(slide, "ABT  |  DSN 2027 research-in-progress", 0.68, 7.14, 8.8, 0.18, 9, MUTED)
    add_text(slide, f"{len(prs.slides):02d}", 12.15, 7.13, 0.5, 0.2, 9, MUTED, align=PP_ALIGN.RIGHT)
    return slide


def add_card(slide, x, y, w, h, title, body, accent=CYAN, body_size=15):
    add_box(slide, x, y, w, h)
    add_box(slide, x, y, 0.055, h, accent, radius=False)
    add_text(slide, title, x + 0.22, y + 0.18, w - 0.42, 0.34, 17, accent, True)
    add_text(slide, body, x + 0.22, y + 0.65, w - 0.42, h - 0.80, body_size, WHITE)


def add_arrow(slide, x1, y1, x2, y2, stroke=CYAN, width=2):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color(stroke)
    shape.line.width = Pt(width)
    shape.line.end_arrowhead = True
    return shape


def add_cell(slide, text, x, y, w, h, size=14, fill=WHITE, bold=False, align=None):
    return add_text(
        slide, text, x, y, w, h, size, fill, bold,
        align=align if align is not None else PP_ALIGN.LEFT,
        valign=MSO_ANCHOR.MIDDLE, margin=0.08,
    )


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "ABT: Goal-Aware Automated Blue Teaming"
    prs.core_properties.subject = "Research-in-progress presentation for DSN 2027"
    prs.core_properties.author = "ABT research team"
    prs.core_properties.keywords = "automated blue teaming, attack graphs, network defense"

    # 1. Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(NAVY)
    add_box(slide, 9.45, 0.0, 3.88, H, PANEL, radius=False)
    add_box(slide, 9.45, 0.0, 0.10, H, CYAN, radius=False)
    add_text(slide, "RESEARCH-IN-PROGRESS  /  DSN 2027", 0.85, 0.88, 7.9, 0.35, 13, CYAN, True)
    add_text(slide, "ABT", 0.82, 1.62, 2.2, 1.0, 52, CYAN, True)
    add_text(
        slide, "Goal-Aware Automated\nBlue Teaming for\nMulti-Host Attack Campaigns",
        0.85, 2.55, 8.1, 2.05, 31, WHITE, True,
    )
    add_text(
        slide, "Simulate candidate network and host defenses before deployment; "
        "adapt to observed attacker behavior.",
        0.88, 4.95, 7.9, 0.82, 18, MUTED,
    )
    add_text(slide, "Authors and affiliations to be finalized", 0.88, 6.65, 6.0, 0.28, 12, MUTED)
    # Abstract network motif on right side
    nodes = [
        (10.15, 1.12, "ATTACKER", RED),
        (11.58, 2.32, "EDGE", AMBER),
        (10.12, 3.68, "CORE", CYAN),
        (11.66, 5.02, "ASSET", GREEN),
        (10.15, 6.0, "ABT", BLUE),
    ]
    for x, y, label, accent in nodes:
        add_box(slide, x, y, 1.22, 0.58, PANEL_2, accent)
        add_text(slide, label, x + 0.06, y + 0.16, 1.1, 0.22, 11, accent, True, align=PP_ALIGN.CENTER)
    add_arrow(slide, 11.28, 1.56, 11.65, 2.22, LINE)
    add_arrow(slide, 11.60, 2.98, 10.80, 3.62, LINE)
    add_arrow(slide, 11.22, 4.26, 11.72, 4.95, LINE)
    add_arrow(slide, 11.72, 5.63, 10.98, 5.96, LINE)

    # 2. Introduction
    slide = add_base(prs, "01 / Introduction", "Autonomous attacks compress the response window",
                     "The research problem is not only detecting an intrusion; it is choosing a safe, effective response across the network.")
    add_card(slide, 0.72, 2.05, 3.75, 2.12, "01  Campaigns are multi-host",
             "LLM-assisted agents plan, execute, and pivot across hosts. A local alert can imply network-wide risk. [3, 4]",
             RED, 16)
    add_card(slide, 4.78, 2.05, 3.75, 2.12, "02  Defenses have side effects",
             "Blocking or isolating can disrupt legitimate workflows, availability, and performance.",
             AMBER, 16)
    add_card(slide, 8.84, 2.05, 3.75, 2.12, "03  Static rules miss context",
             "The best control depends on the target asset, reachable attack paths, and operational constraints.",
             CYAN, 16)
    add_box(slide, 0.72, 4.62, 11.87, 1.48, PANEL_2)
    add_text(slide, "THESIS", 1.02, 4.92, 1.25, 0.28, 12, CYAN, True)
    add_text(slide, "Treat defense as a constrained planning problem: model the attack surface, "
             "simulate candidate actions, then adapt from observed outcomes.",
             2.28, 4.82, 9.72, 0.82, 23, WHITE, True, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, "Evidence so far is a runnable prototype and synthetic-network demonstration—not a validated operational result.",
             0.78, 6.48, 11.7, 0.34, 12, AMBER, True)

    # 3. Gap and questions
    slide = add_base(prs, "01 / Introduction", "Research gap and questions",
                     "The proposed contribution is the closed-loop integration—not a new detector, exploit, or MTD technique.")
    add_card(slide, 0.72, 2.0, 4.0, 3.65, "Gap",
             "Attack-graph methods assess reachability and risk.\n\n"
             "Network controls enforce access and containment.\n\n"
             "Autonomous red-team agents provide repeatable campaigns.\n\n"
             "What is missing here: a shared model that selects among defenses while respecting service constraints and learning from outcomes.",
             BLUE, 15)
    add_card(slide, 5.02, 2.0, 2.4, 1.62, "RQ1 · Security",
             "Does planning reduce goal success versus baselines?", CYAN, 14)
    add_card(slide, 7.68, 2.0, 2.4, 1.62, "RQ2 · QoS",
             "What availability, latency, and performance costs follow?", AMBER, 14)
    add_card(slide, 10.34, 2.0, 2.4, 1.62, "RQ3 · Scale",
             "How does planning scale with hosts, flows, and goals?", GREEN, 14)
    add_card(slide, 5.02, 4.02, 3.58, 1.62, "RQ4 · Adaptation",
             "How quickly does feedback improve the defensive state?", BLUE, 14)
    add_card(slide, 8.88, 4.02, 3.86, 1.62, "RQ5 · Robustness",
             "How sensitive are decisions to uncertain attack and detection probabilities?", RED, 14)

    # 4. Related work
    slide = add_base(prs, "02 / Related Work", "Related work: useful components, open integration problem",
                     "ABT draws on attack-graph assessment, red-team automation, and network-security architecture.")
    x0, y0 = 0.74, 1.98
    cols = [2.15, 3.2, 3.5, 3.0]
    headers = ["Line of work", "Representative work", "What it contributes", "ABT position"]
    xpos = [x0]
    for width in cols[:-1]:
        xpos.append(xpos[-1] + width)
    for x, width, text in zip(xpos, cols, headers):
        add_box(slide, x, y0, width, 0.48, PANEL_2, radius=False)
        add_cell(slide, text, x + 0.05, y0, width - 0.10, 0.48, 12, CYAN, True)
    rows = [
        ("Attack graphs", "HARMer [1]", "Hierarchical attack representation and risk assessment", "Use HARM as planning substrate"),
        ("Asset security", "Enoch et al. [2]", "Value- and requirement-aware security assessment", "Add availability / performance constraints"),
        ("Automated red teams", "Incalmo; APT-Agent [3, 4]", "Multi-host attack planning and execution", "Use adapters; live integration remains future work"),
        ("Benchmarks / network architecture", "MHBench; NIST [5–7]", "Reusable cyber ranges; enterprise and zero-trust network guidance", "Import topologies; model enforceable controls"),
    ]
    rh = 0.90
    for ri, row in enumerate(rows):
        y = y0 + 0.52 + ri * rh
        for ci, (x, width, text) in enumerate(zip(xpos, cols, row)):
            add_box(slide, x, y, width, rh - 0.04, PANEL if ri % 2 == 0 else "101A29", radius=False)
            add_cell(slide, text, x + 0.06, y + 0.03, width - 0.12, rh - 0.10,
                     12 if ci else 13, WHITE if ci else AMBER, ci == 0)
    add_text(slide, "Positioning: combine goal-aware attack-path risk, QoS-constrained action planning, and repeated feedback in one evaluation loop.",
             0.80, 6.20, 11.8, 0.48, 15, WHITE, True)

    # 5. Proposed research and threat model
    slide = add_base(prs, "03 / Proposed Research", "System model and threat assumptions",
                     "Initial scope: one modeled enterprise network, a known asset/goal set, and an adaptive but bounded red-team campaign.")
    add_card(slide, 0.72, 1.98, 3.55, 1.60, "Defender knows",
             "Topology, services, candidate vulnerabilities, asset values, and legitimate flows.",
             CYAN, 14)
    add_card(slide, 0.72, 3.87, 3.55, 1.60, "Attacker can",
             "Enter at modeled entry points, exploit reachable vulnerabilities, pivot, and pursue declared goals.",
             RED, 14)
    add_card(slide, 0.72, 5.76, 3.55, 0.82, "Out of scope initially",
             "Unknown zero-days; production actuation; human deception of attacker.", AMBER, 13)
    # Network graph illustration
    graph = [
        (5.00, 2.25, "Internet", RED),
        (7.25, 2.25, "Web / DMZ", AMBER),
        (9.48, 2.25, "App tier", CYAN),
        (11.48, 4.12, "DB asset", GREEN),
        (8.20, 4.12, "File tier", BLUE),
        (5.00, 4.12, "Identity", BLUE),
    ]
    for x, y, label, accent in graph:
        add_box(slide, x, y, 1.62, 0.70, PANEL_2, accent)
        add_text(slide, label, x + 0.05, y + 0.20, 1.52, 0.25, 13, accent, True, align=PP_ALIGN.CENTER)
    add_arrow(slide, 6.64, 2.60, 7.23, 2.60)
    add_arrow(slide, 8.88, 2.60, 9.45, 2.60)
    add_arrow(slide, 10.84, 2.86, 11.75, 4.05)
    add_arrow(slide, 8.10, 2.94, 8.90, 4.05, LINE)
    add_arrow(slide, 6.55, 2.95, 5.86, 4.05, LINE)
    add_text(slide, "Each edge is permitted connectivity; each destination service contributes exploit risk.",
             4.85, 5.36, 7.5, 0.40, 14, MUTED)
    add_text(slide, "Question to validate: how much topology / exploit-state fidelity is sufficient for safe action ranking?",
             4.85, 5.92, 7.45, 0.58, 17, AMBER, True)

    # 6. Architecture
    slide = add_base(prs, "03 / Proposed Research", "ABT: model → simulate → recommend → observe",
                     "A common network representation lets benchmark topologies and defense actions use the same planner.")
    cards = [
        ("1 · Import", "MHBench / APT topology\n→ Network IR", BLUE),
        ("2 · Model", "Assets + flows + services\n+ vulnerabilities + goals", CYAN),
        ("3 · Analyze", "HARM paths + weighted\ngoal-achievement risk", AMBER),
        ("4 · Plan", "Simulate action sets;\nscore utility and constraints", GREEN),
        ("5 · Observe", "Attack result + MIR-like\ndetection → re-plan", RED),
    ]
    start, width, gap, y = 0.72, 2.18, 0.20, 2.45
    for i, (title, body, accent) in enumerate(cards):
        x = start + i * (width + gap)
        add_card(slide, x, y, width, 1.86, title, body, accent, 14)
        if i < len(cards) - 1:
            add_arrow(slide, x + width + 0.02, y + 0.92, x + width + gap - 0.02, y + 0.92)
    add_box(slide, 1.25, 5.04, 10.85, 0.98, PANEL_2)
    add_text(slide, "Safety boundary", 1.55, 5.36, 1.78, 0.28, 14, AMBER, True)
    add_text(slide, "Prototype mutates a simulated copy only. Real-world enforcement requires a separately validated controller and human approval.",
             3.42, 5.24, 8.20, 0.52, 16, WHITE, True, valign=MSO_ANCHOR.MIDDLE)

    # 7. Network controls, non-MTD
    slide = add_base(prs, "03 / Proposed Research", "Network-level defenses beyond MTD",
                     "The current model has a basic permitted-flow set. The following portfolio is a proposed extension—not yet implemented.")
    cw = [2.30, 3.35, 3.55, 2.30]
    xstarts = [0.72, 3.02, 6.37, 9.92]
    headers = ["Mode / control", "Defensive effect", "Simulator + deployment path", "Selection"]
    for x, w, t in zip(xstarts, cw, headers):
        add_box(slide, x, 1.98, w, 0.44, PANEL_2, radius=False)
        add_cell(slide, t, x + 0.04, 1.98, w - 0.08, 0.44, 12, CYAN, True)
    control_rows = [
        ("PROACTIVE · Network ACL", "Least-privilege segmentation preserves required business flows while cutting lateral paths. [6, 13]", "Planned SDN testbed: Mininet + OVS on Linux; Docker may host endpoints, not the OpenFlow datapath.", "SELECT", CYAN),
        ("REACTIVE · Network block", "Block a specific attacker-relevant flow after a confirmed alert.", "Existing `block_flow` simulation; add typed trigger and narrow firewall adapter.", "SELECT", RED),
        ("PROACTIVE · Host hardening", "Patch a known flaw or disable an explicitly unused exposed service. [14]", "Existing simulation; allowlisted Ansible task on disposable Linux guest.", "SELECT\n(testbed only)", GREEN),
        ("REACTIVE · Host quarantine", "Contain a confirmed compromised host while retaining recovery access. [15]", "Existing isolation is coarse; refine to restricted firewall/EDR profile.", "SELECT\n(refine first)", AMBER),
        ("MTD · Retained", "Software diversity and IP/port shuffle remain distinct comparison controls.", "Existing simulated actions; platform-specific deployment, not removed.", "RETAIN", BLUE),
        ("DEFER · DNS / NTA", "DNS block/sinkhole and network detection need resolver/sensor data. [9, 10, 12]", "Add domain state, telemetry confidence, false positives, and delay.", "OUT OF SCOPE", CYAN),
        ("DEFER · ZTNA / QoS", "Identity-aware access and rate limiting require identity or traffic-load models. [7, 11]", "Add users/devices/sessions or capacity, latency, and throughput first.", "OUT OF SCOPE", MUTED),
    ]
    rh, y0 = 0.53, 2.48
    for ri, (control, effect, model, status, accent) in enumerate(control_rows):
        y = y0 + ri * rh
        bg = PANEL if ri % 2 == 0 else "101A29"
        for x, w in zip(xstarts, cw):
            add_box(slide, x, y, w, rh - 0.035, bg, radius=False)
        add_cell(slide, control, xstarts[0] + 0.08, y, cw[0] - 0.16, rh - 0.05, 9, accent, True)
        add_cell(slide, effect, xstarts[1] + 0.08, y, cw[1] - 0.16, rh - 0.05, 8, WHITE)
        add_cell(slide, model, xstarts[2] + 0.08, y, cw[2] - 0.16, rh - 0.05, 8, WHITE)
        add_cell(slide, status, xstarts[3] + 0.08, y, cw[3] - 0.16, rh - 0.05, 8, accent, True)
    add_text(slide, "MTD remains in the portfolio. Simulator actions are not live controls; SDN validation targets Mininet/OVS on isolated Linux.",
             0.78, 6.37, 11.85, 0.36, 13, MUTED, True)

    # 8. Implementation blueprint
    slide = add_base(prs, "03 / Proposed Research", "Implementation blueprint: simulate first, then validate with Mininet/OVS",
                     "MTD does not require SDN. Current controller is an emulator; Mininet/OVS is the planned Linux testbed. Docker may host endpoints but is not an OpenFlow datapath by itself.")
    blueprint = [
        ("PROACTIVE · Network",
         "Preview least-privilege segmentation and MTD IP remapping against a copied topology; apply only the selected feasible plan.",
         "Measure risk, availability, performance, cost, SDN policy updates and changed virtual-IP mappings.",
         CYAN),
        ("REACTIVE · Network",
         "On a confirmed / injected alert, preview a narrow flow block and replan; controller policy changes remain simulated.",
         "Test alert-to-action trace, legitimate-flow impact, policy-update count, feasibility.",
         AMBER),
        ("PROACTIVE · Host",
         "Use current patch / disable-service simulation actions; testbed actuation only through allowlisted Ansible tasks on disposable Linux hosts.",
         "Check target state, service health, legitimate workflows and rollback; no generic production patching.",
         BLUE),
        ("REACTIVE · Host",
         "Quarantine a confirmed host only after preserving controller / recovery paths. Keep MTD diversity and shuffle as separate baselines.",
         "Refine coarse isolate model; require approval, timed rollback. Compare MTD-only / non-MTD / combined.",
         GREEN),
    ]
    for i, (title, body, verify, accent) in enumerate(blueprint):
        y = 1.97 + i * 1.13
        add_box(slide, 0.78, y, 11.78, 0.95, PANEL)
        add_box(slide, 0.78, y, 0.06, 0.95, accent, radius=False)
        add_text(slide, title, 1.03, y + 0.12, 2.23, 0.30, 15, accent, True)
        add_text(slide, body, 3.22, y + 0.10, 5.93, 0.72, 12, WHITE)
        add_text(slide, verify, 9.28, y + 0.10, 3.02, 0.72, 11, MUTED)
    add_text(slide, "Defer DNS/NTA (resolver / sensor state), ZTNA (identity / policy enforcement), rate limits (traffic/QoS model), and production actuation.",
             0.84, 6.62, 11.8, 0.28, 12, AMBER, True)

    # 9. Planning formulation
    slide = add_base(prs, "03 / Proposed Research", "Constrained defense planning",
                     "For a candidate action set A, evaluate the simulated network N′ = Apply(N, A).")
    add_box(slide, 0.78, 2.02, 7.55, 1.42, PANEL_2)
    add_text(slide, "U(A) = wₛ · (1 − Risk(N′))  +  wₐ · Availability(N′)  +  wₚ · Performance(N′)  −  w꜀ · Cost(A)",
             1.06, 2.38, 7.02, 0.68, 20, WHITE, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_card(slide, 8.70, 2.02, 3.82, 1.42, "Optimize",
             "Greedy marginal-utility selection; compare with stronger baselines.", CYAN, 14)
    add_card(slide, 0.78, 3.80, 3.72, 1.62, "Security objective",
             "Reduce value-weighted probability of attacker goals over modeled paths.", RED, 15)
    add_card(slide, 4.75, 3.80, 3.72, 1.62, "Operational constraints",
             "Preserve minimum legitimate-flow availability and asset performance.", AMBER, 15)
    add_card(slide, 8.70, 3.80, 3.82, 1.62, "Budget and action cost",
             "Respect budget; account for deployment and operational impact.", GREEN, 15)
    add_text(slide, "Critical validation: path independence, probability calibration, utility-weight sensitivity, and monotonicity under coupled actions.",
             0.88, 5.92, 11.6, 0.48, 15, AMBER, True)

    # 10. Learning loop
    slide = add_base(prs, "03 / Proposed Research", "Closed-loop adaptation: current prototype behavior",
                     "Implemented loop: attack → partial observation → re-plan → apply to model → log outcome.")
    flow = [
        ("1", "Attack", "Adapter returns\nsuccess / path / goal", RED),
        ("2", "Detect", "Sample detection on\nintermediate hosts", AMBER),
        ("3", "Update", "Mark detected hosts\nas compromised", CYAN),
        ("4", "Re-plan", "Simulate actions;\nselect feasible plan", BLUE),
        ("5", "Record", "Risk, QoS, cost,\nand success rate", GREEN),
    ]
    for i, (num, title, body, accent) in enumerate(flow):
        x = 0.76 + i * 2.49
        add_box(slide, x, 2.35, 2.15, 1.75, PANEL, accent)
        add_text(slide, num, x + 0.16, 2.51, 0.45, 0.38, 20, accent, True)
        add_text(slide, title, x + 0.65, 2.51, 1.30, 0.34, 16, accent, True)
        add_text(slide, body, x + 0.18, 3.08, 1.78, 0.72, 14, WHITE, align=PP_ALIGN.CENTER)
        if i < len(flow) - 1:
            add_arrow(slide, x + 2.17, 3.18, x + 2.44, 3.18)
    add_card(slide, 0.78, 4.68, 5.63, 1.18, "Implemented",
             "Deterministic seedable loop; AttackResult; configurable detection probability; stability-based stop.",
             GREEN, 14)
    add_card(slide, 6.70, 4.68, 5.77, 1.18, "Not yet established",
             "Online learning quality, real sensor fidelity, adaptive live attacker integration, and production safety.",
             AMBER, 14)
    add_text(slide, "TrainingLoop supports an optional proactive initial plan, then alert-triggered response. SDN model is simulation-only; no real controller is programmed.",
             0.85, 6.23, 11.8, 0.42, 13, MUTED, True)

    # 11. Experimental and verification plan
    slide = add_base(prs, "04 / Experiments & Verification", "Evaluation plan: answer the RQs with controlled comparisons",
                     "Use generated benchmark topologies and synthetic stress tests; keep attack tooling isolated from production networks.")
    x0, y0 = 0.75, 1.98
    cols = [1.45, 3.50, 3.30, 3.42]
    xpos = [x0]
    for c in cols[:-1]:
        xpos.append(xpos[-1] + c)
    for x, w, t in zip(xpos, cols, ["Question", "Experimental design", "Baselines / ablations", "Primary metrics"]):
        add_box(slide, x, y0, w, 0.45, PANEL_2, radius=False)
        add_cell(slide, t, x + 0.05, y0, w - 0.10, 0.45, 12, CYAN, True)
    eval_rows = [
        ("RQ1\nSecurity", "Repeated campaigns across topology families; compare pre/post action outcomes.", "No defense; static ACL; host-only; network-only; full ABT.", "Goal success; residual risk; time-to-containment."),
        ("RQ2\nQoS", "Replay legitimate workload traces under each defense plan.", "Unconstrained planner; per-control ablations.", "Flow availability; latency / throughput; violations."),
        ("RQ3\nScale", "Grow hosts, edges, goals, and candidate actions; profile planner.", "Greedy vs beam / exact on small instances.", "Runtime; memory; utility gap; action count."),
        ("RQ4\nAdapt", "Vary detection rate, attacker persistence, and model error across seeds.", "Static plan; no feedback; perfect / noisy detection.", "Rounds to stable plan; residual risk; variance."),
    ]
    rh = 0.90
    for ri, row in enumerate(eval_rows):
        y = y0 + 0.50 + ri * rh
        for ci, (x, w, txt) in enumerate(zip(xpos, cols, row)):
            add_box(slide, x, y, w, rh - 0.04, PANEL if ri % 2 == 0 else "101A29", radius=False)
            add_cell(slide, txt, x + 0.07, y + 0.02, w - 0.14, rh - 0.08, 11, AMBER if ci == 0 else WHITE, ci == 0)
    add_text(slide, "Protocol: fixed seeds + multiple runs; report confidence intervals; publish configs, importer fixtures, and raw metric outputs.",
             0.83, 6.18, 11.75, 0.48, 14, GREEN, True)

    # 12. Current validation
    slide = add_base(prs, "04 / Experiments & Verification", "What has been verified—and what has not",
                     "Separate functional validation from research-grade empirical evidence.")
    add_card(slide, 0.75, 2.00, 5.75, 3.72, "Verified in the prototype",
             "• HARM paths and goal probability on a synthetic five-host network\n\n"
             "• MHBench- and APT-Agent-shaped importer fixtures\n\n"
             "• Candidate-action simulation and availability constraints\n\n"
             "• Seeded multi-round loop with partial-detection tests\n\n"
             "• Demo: illustrative model risk 0.468 → 0.000; availability 1.00 → 1.00",
             GREEN, 15)
    add_card(slide, 6.78, 2.00, 5.75, 3.72, "Not yet validated",
             "• Real MHBench campaign results or live attacker integration\n\n"
             "• Calibrated exploit / detection probabilities\n\n"
             "• Statistical evidence, confidence intervals, or scalability results\n\n"
             "• Operational network-controller actuation\n\n"
             "• The demo values are simulator outputs, not measured security efficacy",
             AMBER, 15)
    add_text(slide, "The current demo is a feasibility check only; it must not be reported as an experimental result.",
             0.86, 6.13, 11.5, 0.40, 15, RED, True)

    # 13. Limitations
    slide = add_base(prs, "05 / Limitations", "Limitations and threats to validity",
                     "These are research risks to address—not footnotes to hide.")
    limitations = [
        ("Model fidelity", "Imported schemas omit or approximate protocols, identity, state, and application semantics.", BLUE),
        ("Probability calibration", "CVSS and flat playbook defaults are not calibrated exploit-success probabilities.", AMBER),
        ("Dependence assumptions", "Noisy-OR across vulnerabilities and paths may overcount correlated opportunities.", RED),
        ("Detection / adaptation", "Simulated MIR detection and persistence do not reproduce real sensor errors or attacker replanning.", CYAN),
        ("Utility subjectivity", "Weights and hard thresholds can change rankings; conduct sensitivity analysis with stakeholders.", GREEN),
        ("Safety / external validity", "No live enforcement yet; simulator outcomes may not transfer to production networks.", BLUE),
    ]
    for i, (title, body, accent) in enumerate(limitations):
        col = i % 3
        row = i // 3
        add_card(slide, 0.75 + col * 4.06, 2.02 + row * 2.05, 3.74, 1.70, title, body, accent, 13)
    add_text(slide, "Mitigation: calibrate on isolated testbeds, perturb model assumptions, and retain a human approval gate for real actions.",
             0.84, 6.35, 11.7, 0.40, 14, WHITE, True)

    # 14. Roadmap
    slide = add_base(prs, "06 / Conclusion", "From prototype to a defensible DSN contribution",
                     "Prioritize measurable evidence and model validity before expanding the action catalogue.")
    roadmap = [
        ("NOW", "Define typed policy semantics and tests; refine segmentation actions from existing flow primitives.", CYAN),
        ("NEXT", "Add egress / exfiltration semantics, then DNS decisions and typed network telemetry.", BLUE),
        ("THEN", "Add graduated quarantine; calibrate models and run isolated benchmark campaigns.", AMBER),
        ("FINAL", "Compare planners; analyze QoS / uncertainty; release reproducible artifacts.", GREEN),
    ]
    for i, (label, text, accent) in enumerate(roadmap):
        y = 2.06 + i * 0.96
        add_box(slide, 0.84, y, 1.50, 0.64, PANEL_2, accent)
        add_text(slide, label, 0.94, y + 0.19, 1.30, 0.22, 13, accent, True, align=PP_ALIGN.CENTER)
        add_box(slide, 2.58, y, 9.80, 0.64, PANEL)
        add_text(slide, text, 2.82, y + 0.13, 9.30, 0.36, 16, WHITE, valign=MSO_ANCHOR.MIDDLE)
        if i < len(roadmap) - 1:
            add_box(slide, 1.56, y + 0.64, 0.035, 0.28, LINE, radius=False)
    add_text(slide, "Key decision gate: only claim closed-loop effectiveness after repeated, statistically supported benchmark results.",
             0.88, 6.18, 11.7, 0.44, 14, AMBER, True)

    # 15. Conclusion
    slide = add_base(prs, "06 / Conclusion", "Conclusion",
                     "ABT is a promising integration hypothesis; the next contribution is rigorous validation.")
    add_box(slide, 0.86, 2.05, 11.55, 2.45, PANEL_2)
    add_text(slide, "A blue team should ask:", 1.20, 2.43, 3.1, 0.36, 19, CYAN, True)
    add_text(slide, "Which feasible defense action most reduces the attacker’s goal probability—"
             "and does it preserve the services we must protect?",
             1.18, 3.02, 10.55, 0.86, 27, WHITE, True, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, "ABT’s proposal: answer with an attack-graph model, action simulation, operational constraints, and observed feedback.",
             1.20, 4.98, 10.8, 0.62, 17, MUTED)
    add_text(slide, "Thank you  ·  Questions", 1.20, 6.02, 8.0, 0.48, 22, AMBER, True)

    # 16. References
    slide = add_base(prs, "References", "Selected references",
                     "Inline [n] citations map to these sources. Verify bibliographic metadata before submission.")
    refs = [
        "[1] S. Y. Enoch et al., “HARMer: Cyber-attacks Automation and Evaluation,” IEEE Access, 2020. doi:10.1109/ACCESS.2020.3007849",
        "[2] S. Y. Enoch et al., “Novel Security Models, Metrics and Security Assessment for Maritime Vessel Networks,” Computer Networks, 2021. doi:10.1016/j.comnet.2021.107934",
        "[3] “Incalmo: An Autonomous LLM-Assisted System for Red Teaming Multi-Host Networks,” arXiv:2501.16466, 2025.",
        "[4] “APT-Agent: Automated Penetration Testing using Large Language Models,” arXiv:2605.24949, 2026.",
        "[5] MHBench, Multi-Host Cyberrange Environment Generator, github.com/cylabcyberautonomy/MHBench.",
        "[6] R. Chandramouli, NIST SP 800-215, Guide to a Secure Enterprise Network Landscape, 2022. doi:10.6028/NIST.SP.800-215",
        "[7] S. Rose et al., NIST SP 800-207, Zero Trust Architecture, 2020. doi:10.6028/NIST.SP.800-207",
        "[8] MITRE D3FEND, Outbound Traffic Filtering (D3-OTF): https://d3fend.mitre.org/technique/d3f:OutboundTrafficFiltering/",
        "[9] MITRE D3FEND, Network Traffic Analysis (D3-NTA): https://d3fend.mitre.org/technique/d3f:NetworkTrafficAnalysis/",
        "[10] CISA, Protective Domain Name System (DNS) Resolver: https://www.cisa.gov/resources-tools/services/protective-domain-name-system-dns-resolver",
        "[11] CISA, Zero Trust Maturity Model v2.0: https://www.cisa.gov/resources-tools/resources/zero-trust-maturity-model",
        "[12] MITRE D3FEND, DNS Denylisting (D3-DNSDL): https://d3fend.mitre.org/technique/d3f:DNSDenylisting/",
        "[13] K. Scarfone and P. Hoffman, NIST SP 800-41 Rev. 1, Guidelines on Firewalls and Firewall Policy, 2009. doi:10.6028/NIST.SP.800-41r1",
        "[14] NIST SP 800-53 Rev. 5, Security and Privacy Controls, 2020. doi:10.6028/NIST.SP.800-53r5",
        "[15] CISA, Federal Government Cybersecurity Incident and Vulnerability Response Playbooks: cisa.gov/resources-tools/resources/federal-government-cybersecurity-incident-and-vulnerability-response-playbooks",
    ]
    add_bullets(slide, refs[:8], 0.92, 1.88, 5.70, 5.10, 10, WHITE, gap=5)
    add_bullets(slide, refs[8:], 6.78, 1.88, 5.76, 5.10, 10, WHITE, gap=5)

    prs.save(OUT)
    print(f"Wrote {OUT} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    try:
        build()
    except Exception as exc:
        print(f"Failed to build presentation: {exc}", file=sys.stderr)
        raise
